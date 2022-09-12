from baseIncludes import *
import openpyxl

class task_workbook:
  wb = openpyxl.load_workbook(tasks_xls_path)
  shortlist = wb["ShortList"] 
  excel_columns = {
    "count": "A",
    "task": 'C',
    "prio": "D",
    "due_date": "E",
    "class": "F",
    "today": "G",
    "reminder_time": "H",
    "step2": "I",
    "complete_steps": "J",
    "action_date": "K",
    "days_left": "L"
  }
  trace_obj = trace()

  def __init__(self, debug=False):
    trace_obj.prologue(currentframe().f_code.co_name)
    self.debug = debug
    trace_obj.epilogue(currentframe().f_code.co_name)

  def name(self, row_id):
    return self.shortlist[self.excel_columns["task"] + str(row_id)].value

  def prio(self, row_id):
    return self.shortlist[self.excel_columns["prio"] + str(row_id)].value

  def action_date(self, row_id):
    return self.shortlist[self.excel_columns["action_date"] + str(row_id)].value

  def reminder_time(self, row_id):
    return self.shortlist[self.excel_columns["reminder_time"] + str(row_id)].value

  def due_date(self, row_id):
    return self.shortlist[self.excel_columns["due_date"] + str(row_id)].value

  def days_left(self, row_id):
    return self.shortlist[self.excel_columns["days_left"] + str(row_id)].value

  def next_step(self, row_id):
    return self.shortlist[self.excel_columns["step2"] + str(row_id)].value

  def restore_table_format(self):
    tb = self.shortlist.tables["tasksTable"]
    tb.tableStyleInfo.showRowStripes=True
    tb.tableStyleInfo.showColumnStripes=True
  
  def insert_row(self):
    trace_obj.prologue(currentframe().f_code.co_name)
    from copy import copy
    #pdb.set_trace()
    row_id = 2
    task = self.name(row_id)
    while task != None:
      row_id = row_id+1
      task = self.name(row_id)
    self.shortlist.insert_rows(row_id)
    for key in self.excel_columns:
      self.shortlist[self.excel_columns[key] + str(row_id)]._style = copy(self.shortlist[self.excel_columns[key] + str(row_id-1)]._style)
    tb = self.shortlist.tables["tasksTable"]
    tb.ref = "C1:L" + str(row_id)
    tb.tableStyleInfo.showRowStripes=True
    tb.tableStyleInfo.showColumnStripes=True
    self.update_outdated_row(row_id, needs_date_update=True, needs_name_update=True)
    trace_obj.epilogue(currentframe().f_code.co_name)

  def _update_cell(self, column_id, row_id, val, is_date=False):
    trace_obj.prologue(currentframe().f_code.co_name + " : ")
    self.shortlist[column_id +  str(row_id)].value = val
    if is_date == True:
      self.shortlist[column_id +  str(row_id)].number_format = 'mm/dd/yyyy;@'
    self.wb.save(tasks_xls_path)
    trace_obj.epilogue(currentframe().f_code.co_name)

  def update_task_name(self, row_id, task_name):
    trace_obj.prologue(currentframe().f_code.co_name)
    self._update_cell(self.excel_columns["task"], row_id, task_name)
    trace_obj.epilogue(currentframe().f_code.co_name)

  def update_task_prio(self, row_id, task_prio):
    trace_obj.prologue(currentframe().f_code.co_name)
    self._update_cell(self.excel_columns["prio"], row_id, task_prio)
    trace_obj.epilogue(currentframe().f_code.co_name)

  def update_reminder_time_for_task(self, row_id, time):
    trace_obj.prologue(currentframe().f_code.co_name)
    self._update_cell(self.excel_columns["reminder_time"], row_id, time)
    trace_obj.epilogue(currentframe().f_code.co_name)

  def update_action_date_for_task(self, row_id, date):
    trace_obj.prologue(currentframe().f_code.co_name)
    self._update_cell(self.excel_columns["action_date"], row_id, date, True)
    trace_obj.epilogue(currentframe().f_code.co_name)

  def update_action_date_delta_for_task(self, row_id, increment):
    trace_obj.prologue(currentframe().f_code.co_name)
    new_date = datetime.now() + timedelta(days=int(increment))
    self.update_action_date_for_task(row_id, new_date.date())
    if increment != 0:
      new_remind_time = datetime.strptime("06:00 AM", "%I:%M %p").time()
      self.update_reminder_time_for_task(row_id, new_remind_time) 
    trace_obj.epilogue(currentframe().f_code.co_name)

  def update_due_month_delta_for_task(self, row_id, increment):
    #also updates action date as it doesn't make sense to leave action date behind!
    trace_obj.prologue(currentframe().f_code.co_name)
    curr_due_date = self.due_date(row_id)
    new_due_date = curr_due_date + timedelta(days=int(increment*30))
    self._update_cell(self.excel_columns["due_date"], row_id, new_due_date, True)
    self._update_cell(self.excel_columns["action_date"], row_id, new_due_date, True)
    new_remind_time = datetime.strptime("06:00 AM", "%I:%M %p").time()
    self.update_reminder_time_for_task(row_id, new_remind_time) 
    trace_obj.epilogue(currentframe().f_code.co_name)

  def update_reminder_time_delta_for_task(self, row_id, increment):
    trace_obj.prologue(currentframe().f_code.co_name)
    #trace_obj.trace("row: " + str(row_id) + " increment: " + str(increment))
    self.update_reminder_time_for_task(row_id, (datetime.now() +\
      timedelta(minutes=increment)).time())
    trace_obj.epilogue(currentframe().f_code.co_name)

  def update_outdated_row(self, row_id=2, needs_date_update=True,\
    needs_name_update=False):
    trace_obj.prologue(currentframe().f_code.co_name)
    #pdb.set_trace()
    if needs_name_update == True:
      task = get_input_from_user("input TASK name: ")
      self.update_task_name(row_id, task)
      prio = get_input_from_user("input Prio (0:now 1:Plan for it 2:good to do: ")
      self.update_task_prio(row_id, prio)
    else:
      task = self.name(row_id)
    if needs_date_update == True:
      delta = get_input_from_user("DAYS to increment from today for " + task + ": ")
      self.update_action_date_delta_for_task(row_id, int(delta))
    if int(delta) == 0:   
      date_input = input("TIME as HH:MM PM/AM: ")
      if date_input == "0":
        date_input = datetime.now()
        remind_time = date_input.time()
      else:
        remind_time = datetime.strptime(date_input, "%I:%M %p").time()
    else:
      remind_time = datetime.strptime("06:00 AM", "%I:%M %p").time()
    self.update_reminder_time_for_task(row_id, remind_time)
    if needs_name_update == True:
      self._update_cell(self.excel_columns["count"], row_id,\
        "==INDIRECT(\"R[-1]C\",0)+1")
      self._update_cell(self.excel_columns["days_left"],\
        row_id, "=" + self.excel_columns["action_date"] + str(row_id) + "-TODAY()")
      next_step = input("input next step")
      self._update_cell(self.excel_columns["step2"], row_id, next_step)
    trace_obj.epilogue(currentframe().f_code.co_name)

